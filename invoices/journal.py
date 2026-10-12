"""One place that answers "where is this invoice?": every invoice the tool ever saw, with
its status. journal.json is the record; יומן חשבוניות.xlsx is only a view rebuilt from it,
so a log she left open in Excel loses nothing - the view just catches up next time."""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .workbook import GREEN, HEAD, YELLOW, _no_formulas

PENDING = "ממתינה לבדיקה"
READY = "מוכנה לייצוא"
EXPORTED = "בקובץ קליטה"     # not "נקלטה": the tool cannot see whether she imported it
REJECTED = "לא תיקלט"

ORDER = {PENDING: 0, READY: 1, EXPORTED: 2, REJECTED: 3}
FILLS = {PENDING: YELLOW, READY: GREEN, EXPORTED: PatternFill("solid", fgColor="DDEBF7"),
         REJECTED: PatternFill("solid", fgColor="EDEDED")}
COLS = [("status", "מצב"), ("supplier_name", "ספק"), ("invoice_number", "מספר חשבונית"),
        ("invoice_date", "תאריך חשבונית"), ("total", 'סה"כ'), ("todo", "מה צריך / למה"),
        ("updated", "עודכן"), ("export_file", "קובץ קליטה")]

JSON_NAME = "journal.json"
XLSX_NAME = "יומן חשבוניות.xlsx"
LOCKED = "\n\nיומן החשבוניות פתוח באקסל ולכן לא התעדכן הפעם. הוא יתעדכן בפעם הבאה שהוא סגור."


def status_after_fetch(row: dict) -> tuple[str, str]:
    answer = str(row.get("approve") or "").strip()
    if answer == "לא":
        return REJECTED, row.get("notes") or 'סומנה "לא"'
    if answer == "כן":           # she (or the tool, for a clean row) already approved it
        return READY, ""
    if row.get("notes"):
        return PENDING, row["notes"]
    return READY, ""


def load(home: Path) -> dict:
    path = home / JSON_NAME
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def _as_date(v):
    if isinstance(v, (date, datetime)) or not v:
        return v
    try:
        return datetime.fromisoformat(str(v)).date()
    except ValueError:
        return v


def record(home: Path, changes: list[tuple[dict, str, str]], export_file: str = "") -> str:
    """changes: (review row, status, reason). Saves the record and rebuilds the view.
    Returns a line to add to her message ("" when all went well)."""
    entries = load(home)
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for row, status, reason in changes:
        entries[row["key"]] = {
            "status": status,
            # a scan has no readable name: the file name is how she finds it
            "supplier_name": row.get("supplier_name") or (f'קובץ: {Path(row["file"]).name}' if row.get("file") else ""),
            "invoice_number": row.get("invoice_number") or "",
            "invoice_date": row.get("invoice_date"),
            "total": row.get("total"),
            "todo": reason,
            "updated": now,
            "export_file": export_file if status == EXPORTED else "",
        }
    tmp = home / (JSON_NAME + ".tmp")
    tmp.write_text(json.dumps(entries, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    tmp.replace(home / JSON_NAME)
    try:
        write_view(entries, home / XLSX_NAME)
    except PermissionError:
        return LOCKED
    return ""


def write_view(entries: dict, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "יומן"
    ws.sheet_view.rightToLeft = True
    ws.append([title for _, title in COLS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    # open work first, then the newest history
    items = sorted(entries.values(), key=lambda e: e["updated"], reverse=True)
    items.sort(key=lambda e: ORDER.get(e["status"], 9))
    for e in items:
        ws.append([_as_date(e.get(k)) if k == "invoice_date" else e.get(k, "") for k, _ in COLS])
        for cell in ws[ws.max_row]:
            cell.fill = FILLS.get(e["status"], YELLOW)
        ws.cell(ws.max_row, 4).number_format = "DD/MM/YYYY"
        ws.cell(ws.max_row, 5).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGH", [16, 26, 16, 14, 12, 50, 17, 34]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(ws.max_row, 1)}"
    _no_formulas(ws)
    wb.save(path)
