"""The supplier index exported from Hashavshevet (אינדקס ספקים) as Excel.

With it, a supplier seen for the first time already gets its supplier card, matched by
ח.פ / ע.מ - never by name: two cards can share a name, and names are spelled loosely.
The export has title lines above the real header; the header row is the one holding
"מפתח חשבון" and a "מורשה" column ("מס' ע.מורשה").
"""
from __future__ import annotations

import csv
import re
import tempfile
from pathlib import Path


def _rows_openpyxl(path: Path) -> list[list]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _rows_excel(path: Path) -> list[list]:
    """Hashavshevet's own Excel exports often break openpyxl (style errors): let the
    installed Excel save the sheet as CSV and read that."""
    import pythoncom
    import win32com.client
    out = Path(tempfile.mkdtemp()) / "index.csv"
    pythoncom.CoInitialize()
    excel = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")  # private instance, never her open Excel
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(str(path.resolve()), 0, True)
        wb.SaveAs(str(out), FileFormat=62)  # 62 = xlCSVUTF8
        wb.Close(False)
    finally:
        if excel is not None:
            excel.Quit()
        pythoncom.CoUninitialize()
    with open(out, encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse(rows: list[list]) -> dict[str, list[dict]]:
    """ח.פ -> every card that carries it (usually one)."""
    out: dict[str, list[dict]] = {}
    cols = None
    for row in rows:
        cells = [_text(c) for c in row]
        if cols is None:
            key = next((i for i, c in enumerate(cells) if c == "מפתח חשבון"), None)
            vid = next((i for i, c in enumerate(cells) if "מורשה" in c), None)
            if key is not None and vid is not None:
                name = next((i for i, c in enumerate(cells) if c.startswith("שם")), None)
                cols = (key, name, vid)
            continue
        key, name, vid = cols
        get = lambda i: cells[i] if i is not None and i < len(cells) else ""
        digits = re.sub(r"\D", "", get(vid))
        account = get(key)
        if not digits or not account:
            continue
        out.setdefault(digits.zfill(9), []).append({"supplier_account": account, "name": get(name)})
    return out


def load(path: str | Path) -> dict[str, list[dict]]:
    """Empty when no index is configured. Errors propagate: a configured index that
    cannot be read must be reported, not silently ignored."""
    if not path:
        return {}
    return parse(read_rows(Path(path)))


def read_rows(path: Path) -> list[list]:
    """Any Hashavshevet Excel export as rows: openpyxl first, the installed Excel if it chokes."""
    try:
        return _rows_openpyxl(path)
    except Exception:
        return _rows_excel(path)
