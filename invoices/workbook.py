"""The two Excel files the bookkeeper actually touches.

1. review.xlsx  - one row per invoice. Green rows are ready; yellow rows need her
                  (new supplier, missing field, VAT mismatch). She fills the gaps.
2. import.xlsx  - built from the approved rows, in the shape Hashavshevet ERP's
                  "יבוא תנועות הנה"ח > יבוא מטבלת ODBC" expects: one header row,
                  account keys and references stored as TEXT, and a named range
                  over the data (without a name the import window shows NONAME).
"""
from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

from .extract import journal_reference

REVIEW_COLS = [
    ("approve", "לאשר? (כן/לא)"),
    ("supplier_name", "ספק"),
    ("supplier_id", 'ח.פ / ע.מ'),
    ("supplier_account", "כרטיס ספק"),
    ("expense_account", "כרטיס הוצאה"),
    ("invoice_number", "מס' חשבונית"),
    ("invoice_date", "תאריך"),
    ("net", 'לפני מע"מ'),
    ("vat", 'מע"מ'),
    ("total", 'סה"כ'),
    ("allocation_number", "מספר הקצאה"),
    ("details", "פרטים"),
    ("notes", "הערות"),
    ("file", "קובץ"),
    ("key", "מזהה"),
    ("sender", "נשלח מ"),
]
COL = {k: i for i, (k, _) in enumerate(REVIEW_COLS)}

# Column titles in import.xlsx: Hashavshevet's own field names, so the one-time mapping
# ("שידוך") in the import window is obvious; the mapping is then saved and reused.
IMPORT_COLS = ["סוג תנועה", "אסמכתא 1", "אסמכתא 2", "תאריך אסמכתא", "תאריך ערך", "פרטים",
               "חשבון חובה 1", "סכום חובה 1", "חשבון חובה 2", "סכום חובה 2",
               "חשבון זכות 1", "סכום זכות 1", "חשבון זכות 2", "סכום זכות 2",
               "מספר עמ", "מספר הקצאה"]
TEXT_COLS = (0, 1, 2, 6, 8, 10, 12, 14, 15)   # keys & references as text, per the H-ERP guide


def details_text(invoice_date, description: str, supplier_name: str) -> str:
    """פרטים: "MM/YY <מה נעשה> <ספק>" - month/year first, always a 2-digit month."""
    month = f"{invoice_date:%m/%y}" if invoice_date else ""
    return " ".join(p for p in (month, (description or "").strip(), (supplier_name or "").strip()) if p)[:50]


def description_of(details: str, supplier_name: str) -> str:
    """The "<מה נעשה>" part of her פרטים, to reuse for the same supplier next month."""
    d = re.sub(r"^\s*\d{1,2}/\d{2,4}\s*", "", details or "").strip()
    name = (supplier_name or "").strip()
    if name and d.endswith(name):
        d = d[: -len(name)].strip()
    return d

GREEN = PatternFill("solid", fgColor="E2F0D9")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="1F4E78")


def _no_formulas(ws) -> None:
    """Text read out of a PDF is data, never a formula. openpyxl stores any string that
    starts with "=" as a live formula, so a crafted invoice could plant one in her sheet."""
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"


def write_review(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "לבדיקה"
    ws.sheet_view.rightToLeft = True
    ws.append([title for _, title in REVIEW_COLS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r in rows:
        ws.append([r.get(k, "") for k, _ in REVIEW_COLS])
        fill = YELLOW if r.get("notes") else GREEN
        for cell in ws[ws.max_row]:
            cell.fill = fill
    widths = [12, 26, 13, 12, 12, 14, 12, 12, 10, 12, 14, 40, 45, 30, 8, 28]
    for col, w in zip("ABCDEFGHIJKLMNOP", widths):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[COL["invoice_date"]].number_format = "DD/MM/YYYY"
        for c in row[COL["net"]:COL["total"] + 1]:
            c.number_format = "#,##0.00"
    ws.column_dimensions["O"].hidden = True   # מזהה
    ws.freeze_panes = "B2"
    _no_formulas(ws)
    wb.save(path)


def read_review(path: Path) -> list[dict]:
    ws = load_workbook(path, data_only=True).active
    # by header title, not position: a sheet made by an older version (fewer columns)
    # must still be read correctly after an update
    by_title = {title: k for k, title in REVIEW_COLS}
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    keys = [by_title.get(str(h).strip() if h is not None else "", f"_extra{i}") for i, h in enumerate(header)]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(values):
            continue
        row = {k: None for k, _ in REVIEW_COLS} | dict(zip(keys, values))
        for k in ("supplier_id", "supplier_account", "expense_account", "invoice_number",
                  "allocation_number", "supplier_name", "key", "sender", "details"):
            v = row.get(k)
            # Excel turns typed account numbers into floats (40012 -> 40012.0)
            row[k] = "" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v).strip())
        if not row["key"]:  # a row she typed in herself: give it a stable identity
            row["key"] = f'manual:{row["supplier_id"]}:{row["invoice_number"]}:{row.get("total")}'
        rows.append(row)
    return rows


def _d(v) -> Decimal:
    return Decimal(str(v or 0)).quantize(Decimal("0.01"))


def write_import(rows: list[dict], vat_account: str, path: Path, ttype: str = "") -> None:
    """One journal line per invoice: expense + input VAT debited, supplier credited.

    A credit note (negative total) is the mirror image: supplier debited, expense and
    VAT credited, all amounts positive - safer than handing the import negative sums.
    ttype is the company's ס.ת for supplier invoices (it differs between companies);
    an invoice without VAT (exempt dealer, receipt, municipal tax) gets an empty ס.ת.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "תנועות"
    ws.append(IMPORT_COLS)
    for r in rows:
        net, vat, total = _d(r["net"]), _d(r["vat"]), _d(r["total"])
        ref = journal_reference(r["invoice_number"])  # אסמכתא: no prefix, numeric, max 9
        d = r["invoice_date"]
        d = d.date() if hasattr(d, "date") else d
        head = [ttype if vat else "", ref, ref, d, d]
        tail = [r.get("supplier_id") or "", r.get("allocation_number") or ""]
        vat_acc = vat_account if vat else ""
        # Unused amount cells hold 0, never blank: ODBC guesses each column's type from its
        # first rows, and a column that starts out blank reads a later value as NULL.
        name = r.get("supplier_name") or ""
        if total >= 0:
            details = r.get("details") or details_text(d, "", name)
            ws.append(head + [details[:50],
                              r["expense_account"], float(net), vat_acc, float(vat),
                              r["supplier_account"], float(total), "", 0.0] + tail)
        else:
            details = r.get("details") or details_text(d, "זיכוי", name)
            ws.append(head + [details[:50],
                              r["supplier_account"], float(-total), "", 0.0,
                              r["expense_account"], float(-net), vat_acc, float(-vat)] + tail)
    last = ws.max_row
    for row in ws.iter_rows(min_row=2):
        for idx in TEXT_COLS:
            row[idx].number_format = "@"
            if row[idx].value is not None:
                row[idx].value = str(row[idx].value)
        row[3].number_format = row[4].number_format = "DD/MM/YYYY"
    _no_formulas(ws)
    ref = f"'{ws.title}'!$A$1:$P${last}"
    wb.defined_names["invoices"] = DefinedName("invoices", attr_text=ref)
    wb.save(path)


def save_xls_copy(xlsx: Path) -> Path | None:
    """Also save an Excel 97-2003 .xls next to the .xlsx, through the installed Excel.

    A 32-bit Hashavshevet often reads only .xls through ODBC (the H-ERP guide's own
    example is an .xls). Returns None when Excel is not available.
    """
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return None
    xls = xlsx.with_suffix(".xls")
    pythoncom.CoInitialize()
    excel = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")  # private instance, never her open Excel
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(str(xlsx.resolve()))
        wb.SaveAs(str(xls.resolve()), FileFormat=56)  # 56 = xlExcel8
        wb.Close(False)
        return xls
    except Exception:
        return None
    finally:
        if excel is not None:
            excel.Quit()
        pythoncom.CoUninitialize()
