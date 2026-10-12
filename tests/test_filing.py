"""Where an export puts things, as she asked: a dated folder with every invoice as
"<ספק> <מספר>.pdf", the import file in REP, a field mapping (WPR) beside it, and each
mail moved to "טופל" only when everything in it is handled."""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

import fitz
import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from invoices import journal, main, sources, wpr  # noqa: E402
from invoices.workbook import IMPORT_COLS  # noqa: E402
from test_workflow import fill, home, standard_pass  # noqa: E402,F401  (fixture)

# The field mapping of the .wpr saved at her office, used for a real import (מנה 762).
# Only Hashavshevet's field names: the part that held her folder path is not in here.
REAL_MAPPING = bytes.fromhex(
    "000357494e250000000009f1e5e220faf0e5f2e408e0f1eeebfae0203108e0f1eeebfae020320cfae0f8e9ea20e0f1eeebfae0"
    "09fae0f8e9ea20f2f8ea09f7e5e320faeee7e9f808f7e5e320eee8e1f205f4f8e8e9ed0ce7f9e1e5ef20e7e5e1e420310ce7f9"
    "e1e5ef20e7e5e1e420320ce7f9e1e5ef20e6ebe5fa20310ce7f9e1e5ef20e6ebe5fa20320bf1ebe5ed20e7e5e1e420310bf1eb"
    "e5ed20e7e5e1e420320bf1ebe5ed20e6ebe5fa20310bf1ebe5ed20e6ebe5fa2032000000000000000000000007eef1f4f820f2"
    "ee00000000000aeef1f4f820e4f7f6e0e40000000130013001300130")


def test_mapping_is_the_one_that_worked_at_her_office():
    assert wpr.mapping() == REAL_MAPPING


def test_every_mapped_column_is_in_the_import_file():
    """The mapping feeds each field from a column by name; a missing column fails the import."""
    assert {s for s in wpr.SLOTS if s} <= set(IMPORT_COLS)


@pytest.fixture
def short_dir():
    """The mapping holds the file's path twice in 255 bytes, more than pytest's own temp
    folders leave room for (her real REP path comes to 222 bytes)."""
    import os
    import shutil
    import tempfile
    try:
        d = Path(tempfile.mkdtemp(prefix="hv", dir=os.path.abspath(os.sep)))
    except OSError:
        pytest.skip("no short folder available on this machine")
    yield d
    shutil.rmtree(d, ignore_errors=True)


def test_wpr_points_at_the_file_with_the_right_driver(short_dir):
    b = wpr.write(short_dir / "101500.xls", "305").read_bytes()
    assert b.startswith(b"\x02V2") and b.endswith(REAL_MAPPING)
    text = b.decode("cp1255", errors="replace")
    assert f"DBQ={short_dir / '101500.xls'};DefaultDir={short_dir};DriverId=790;" in text
    assert "\x07Journal\x0b[jurnal305]" in text
    assert b"DriverId=1046" in wpr.build(short_dir / "a.xlsx", "305")
    with pytest.raises(ValueError):                         # 255 bytes is the format's limit
        wpr.build(short_dir / ("x" * 200) / "a.xls", "305")


@pytest.fixture
def client(home, short_dir):
    tmp_path = short_dir
    cfg = main.load(main.CONFIG, {})
    cfg.update(hid="305", files_folder=str(tmp_path / "F" / "לקוח"), rep_folder=str(tmp_path / "REP" / "11-לקוח"))
    main.save(main.CONFIG, cfg)
    return tmp_path


def test_export_files_everything_where_she_keeps_it(client, home):
    fill(main.fetch(), standard_pass)
    out = main.export()
    now = datetime.now()
    rep_day = client / "REP" / "11-לקוח" / "חשבוניות שנקלטו" / f"{now:%d-%m-%Y}"
    day = client / "F" / "לקוח" / f"{now:%d.%m.%y}"
    assert out.parent == rep_day
    wb = openpyxl.load_workbook(out)
    assert wb.active.title == "Journal" and "jurnal305" in wb.defined_names
    assert [c.value for c in wb.active[1]] == IMPORT_COLS
    (mapping,) = rep_day.glob("Journal.*.jurnal305.wpr")
    data_file = out.with_suffix(".xls") if out.with_suffix(".xls").exists() else out   # .xls when Excel is there
    assert mapping.name == f"Journal.{data_file.name}.jurnal305.wpr"
    assert f"DBQ={data_file};DefaultDir={rep_day};" in mapping.read_bytes().decode("cp1255")
    pdfs = sorted(p.name for p in day.glob("*.pdf"))
    assert "חשמל השרון בע״מ 20261045.pdf" in pdfs and len(pdfs) == 5
    assert list(day.glob("[0-9]*.xls*"))                   # a copy of the import file with the invoices


def test_an_invoice_file_already_there_is_never_overwritten(tmp_path):
    a, b = tmp_path / "a.pdf", tmp_path / "b.pdf"
    a.write_bytes(b"one")
    b.write_bytes(b"two")
    dest = tmp_path / "day" / "ספק 1.pdf"
    dest.parent.mkdir()
    assert main._copy_once(a, dest) == dest
    assert main._copy_once(a, dest) == dest                        # the same file again: no copy
    assert main._copy_once(b, dest).name == "ספק 1 (2).pdf"         # another file, same name
    assert dest.read_bytes() == b"one"


def test_without_client_folders_it_works_as_before(home):
    fill(main.fetch(), standard_pass)
    out = main.export()
    assert out.parent == home and "invoices" in openpyxl.load_workbook(out).defined_names
    assert not list(home.glob("*.wpr"))


def test_a_cut_up_file_keeps_its_original_under_makor(client, home):
    from make_demo_invoices import OUT
    doc = fitz.open()
    for name in ("01_electric.pdf", "02_office_table.pdf"):
        with fitz.open(OUT / name) as src:
            doc.insert_pdf(src)
        (OUT / name).unlink()
    doc.save(OUT / "monthly.pdf")
    fill(main.fetch(), standard_pass)
    main.export()
    day = client / "F" / "לקוח" / f"{datetime.now():%d.%m.%y}"
    assert [p.name for p in (day / "מקור").iterdir()] == ["monthly.pdf"]


def test_missing_drive_stops_before_anything_is_written(client, home):
    blocker = client / "file.txt"
    blocker.write_text("x")
    cfg = main.load(main.CONFIG, {})
    cfg["files_folder"] = str(blocker / "under a file")
    main.save(main.CONFIG, cfg)
    fill(main.fetch(), standard_pass)
    with pytest.raises(SystemExit):
        main.export()
    assert not (client / "REP").exists() or not list((client / "REP").rglob("*.xlsx"))
    assert not main.load(main.STATE, {}).get("invoice_keys")      # nothing counted as entered


# --- moving mail to "טופל" ------------------------------------------------------------
T1, T2, T3, T4 = "a" * 12, "b" * 12, "c" * 12, "d" * 12
E, P, R = journal.EXPORTED, journal.PENDING, journal.REJECTED


def entries(*rows):
    return {f"outlook:{tag}-{n}": {"status": s, "todo": todo} for tag, n, s, todo in rows}


@pytest.mark.parametrize("rows, moved", [
    ([(T1, 1, E, ""), (T1, 2, E, "")], [T1]),                         # all in the import file
    ([(T1, 1, E, ""), (T1, 2, P, "חסר")], []),                         # one still waiting
    ([(T1, 1, E, ""), (T1, 2, R, 'סומנה "לא"')], [T1]),                 # the rest turned down
    ([(T2, 1, R, main.DEAL_NOTE)], []),                                # only a חשבון עסקה: stays
    ([(T3, 1, R, main.DUPLICATE_DONE)], [T3]),                         # a duplicate: still moves
    ([(T4, 1, R, 'סומנה "לא"')], []),                                  # nothing real in it
])
def test_her_rule_for_moving_mail(rows, moved):
    assert main.mails_to_move(entries(*rows), {r[0] for r in rows}) == moved


def test_tag_of_a_cut_up_invoice_is_its_mail():
    assert sources.tag_of(f"outlook:{T1}-2#3-4") == T1 and sources.tag_of("file:abc") == ""


def test_export_moves_handled_mail(home, monkeypatch):
    cfg = {**main.load(main.CONFIG, {}), "source": "outlook", "done_folder": "לקוח\\טופל"}
    main.save(main.CONFIG, cfg)
    main.save(main.STATE, {"mail_ids": [], "invoice_keys": [], "entry_ids": {T1: "EID1", T2: "EID2"}})
    journal.record(home, [({"key": f"outlook:{T1}-1"}, E, ""), ({"key": f"outlook:{T2}-1"}, P, "")])
    calls = []
    monkeypatch.setattr(sources, "move_mails", lambda ids, folder: calls.append((ids, folder)) or (len(ids), []))
    text = main._move_handled(cfg, main.load(main.STATE, {}), [{"key": f"outlook:{T1}-1"}, {"key": f"outlook:{T2}-1"}])
    assert calls == [(["EID1"], "לקוח\\טופל")] and "1 מיילים" in text


def test_keys_and_references_are_text_in_the_import_file(home):
    """ODBC reads a number-looking account or allocation number as a number and drops zeros."""
    from invoices.workbook import TEXT_COLS
    fill(main.fetch(), standard_pass)
    ws = openpyxl.load_workbook(main.export()).active
    names = {IMPORT_COLS[i] for i in TEXT_COLS}
    assert {"חשבון חובה 1", "חשבון זכות 1", "מספר עמ", "מספר הקצאה", "אסמכתא 1"} <= names
    for row in ws.iter_rows(min_row=2):
        for i in TEXT_COLS:
            assert row[i].number_format == "@" and (row[i].value is None or isinstance(row[i].value, str))


def test_fetch_remembers_each_mail_so_it_can_be_moved_later(home, monkeypatch):
    from test_outlook import Folder, Mail, install_fake_outlook
    install_fake_outlook(monkeypatch, Folder([], sub={"לקוח": Folder([Mail(31, "x", ["a.pdf"])])}))
    main.save(main.CONFIG, {**main.load(main.CONFIG, {}), "source": "outlook", "outlook_folder": "לקוח"})
    main.fetch()
    eid = f"{31:0>40}"
    assert main.load(main.STATE, {})["entry_ids"] == {sources.mail_tag(eid): eid}
