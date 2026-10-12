"""The log: every invoice she ever got, with where it stands, in one sheet. It may lag
behind (open in Excel), but it must never be wrong and never block the real work."""
from __future__ import annotations

import os
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "demo"))

import make_demo_invoices as demo  # noqa: E402
from invoices import journal, main  # noqa: E402
from invoices.workbook import COL  # noqa: E402


@pytest.fixture
def home(tmp_path, monkeypatch):
    demo.OUT = tmp_path / "mail"
    demo.main()
    h = tmp_path / "home"
    h.mkdir()
    for name, value in (("HOME", h), ("CONFIG", h / "config.json"),
                        ("SUPPLIERS", h / "suppliers.json"), ("STATE", h / "state.json")):
        monkeypatch.setattr(main, name, value)
    main.save(main.CONFIG, {**main.DEFAULT_CONFIG, "source": "folder", "test_folder": str(demo.OUT),
                            "own_company_ids": [demo.OWN_ID], "vat_account": "77100", "ttype": "ח"})
    return h


def view(home: Path) -> list[dict]:
    ws = openpyxl.load_workbook(home / journal.XLSX_NAME).active
    keys = [k for k, _ in journal.COLS]
    assert [c.value for c in ws[1]] == [t for _, t in journal.COLS]
    return [dict(zip(keys, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def mark(review: Path, answers: dict[str, tuple]) -> None:
    """answers: file name -> (approve, supplier_account, expense_account)."""
    wb = openpyxl.load_workbook(review)
    for row in wb.active.iter_rows(min_row=2):
        name = Path(row[COL["file"]].value).name
        if name in answers:
            row[COL["approve"]].value, row[COL["supplier_account"]].value, row[COL["expense_account"]].value \
                = answers[name]
    wb.save(review)


def test_fetch_logs_every_invoice_as_its_review_row_says(home):
    review = main.fetch()
    rows = view(home)
    sheet = list(openpyxl.load_workbook(review).active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(sheet)                      # nothing missing, nothing extra
    for r in rows:                                      # all suppliers are new: every row needs her
        assert r["status"] in (journal.PENDING, journal.REJECTED) and r["todo"]


def test_export_moves_rows_to_their_real_status(home, capsys):
    review = main.fetch()
    mark(review, {"01_electric.pdf": ("כן", "40010", "81100"),       # goes into the import file
                  "02_office_table.pdf": ("כן", "", ""),              # approved but incomplete: held
                  "03_exempt_dealer.pdf": ("לא", "", "")})            # rejected
    out = main.export()
    rows = view(home)
    exported = [r for r in rows if r["status"] == journal.EXPORTED]
    assert [(r["invoice_number"], r["export_file"]) for r in exported] == [("20261045", out.name)]
    forwarded = [r for r in rows if r["invoice_number"] == "20261045" and r is not exported[0]]
    assert forwarded[0]["status"] == journal.PENDING and "כפילות" in forwarded[0]["todo"]
    by_supplier = {r["supplier_name"]: r for r in rows}
    office, exempt = by_supplier['אופיס סנטר ישיר בע"מ'], by_supplier["יעל ברק - עיצוב גרפי"]
    assert office["status"] == journal.PENDING and "חסר כרטיס ספק" in office["todo"]
    assert exempt["status"] == journal.REJECTED
    assert not office["export_file"] and not exempt["export_file"]


def test_open_work_is_listed_before_history(home):
    review = main.fetch()
    mark(review, {"01_electric.pdf": ("כן", "40010", "81100")})
    main.export()
    order = [journal.ORDER[r["status"]] for r in view(home)]
    assert order == sorted(order)
    assert order[0] == journal.ORDER[journal.PENDING] and journal.ORDER[journal.EXPORTED] in order


def test_log_open_in_excel_does_not_block_export_and_catches_up(home, monkeypatch, capsys):
    review = main.fetch()
    mark(review, {"01_electric.pdf": ("כן", "40010", "81100")})
    real = journal.write_view

    def locked(entries, path):
        raise PermissionError("open in Excel")
    monkeypatch.setattr(journal, "write_view", locked)
    capsys.readouterr()
    assert main.export() is not None
    assert "יומן החשבוניות פתוח באקסל" in capsys.readouterr().out
    monkeypatch.setattr(journal, "write_view", real)
    main.run("journal")                                 # next time it is closed: the view catches up
    assert journal.EXPORTED in {r["status"] for r in view(home)}


def test_broken_log_never_turns_an_export_into_a_double_entry(home, monkeypatch, capsys):
    review = main.fetch()
    mark(review, {"01_electric.pdf": ("כן", "40010", "81100")})

    def broken(*a, **k):
        raise RuntimeError("disk full")
    monkeypatch.setattr(journal, "record", broken)
    capsys.readouterr()
    assert main.export() is not None                    # the import file is still made
    out = capsys.readouterr().out
    assert "נוצר קובץ קליטה" in out and "לא התעדכן" in out
    assert main.export() is None                        # and the same invoice is not exported again


def test_journal_command_before_any_fetch_says_it_is_empty(home, capsys):
    main.run("journal")
    assert "היומן עוד ריק" in capsys.readouterr().out
    assert not (home / journal.XLSX_NAME).exists()


def test_pdf_text_cannot_plant_a_formula_in_the_log(home):
    journal.record(home, [({"key": "k", "supplier_name": "=HYPERLINK(\"http://x\")"}, journal.PENDING, "")])
    cell = openpyxl.load_workbook(home / journal.XLSX_NAME).active.cell(2, 2)
    assert cell.data_type == "s"


@pytest.mark.parametrize("name", [main.JOURNAL_BAT, main.APP_BAT])
def test_new_button_reaches_a_pc_whose_updater_did_not_know_it(tmp_path, name):
    app = tmp_path / "app"
    app.mkdir()
    (app / "run.py").write_text("", encoding="utf-8")
    main.ensure_launchers(app)
    # identical to the button shipped in the repo, CRLF included (cmd reads it line by line)
    assert (app / name).read_bytes() == (ROOT / name).read_bytes()


@pytest.mark.parametrize("name", [main.JOURNAL_BAT, main.APP_BAT])
def test_old_pythonw_only_button_is_replaced(tmp_path, name):
    """On her PC "python" is the Microsoft Store shortcut and pythonw does not exist."""
    app = tmp_path / "app"
    app.mkdir()
    (app / "run.py").write_text("", encoding="utf-8")
    (app / name).write_bytes(b'@echo off\r\ncd /d "%~dp0"\r\nstart "" pythonw run.py app\r\n')
    main.ensure_launchers(app)
    assert (app / name).read_bytes() == (ROOT / name).read_bytes()


@pytest.mark.parametrize("name", [n for n in os.listdir(ROOT) if n.endswith(".bat")])
def test_every_button_tries_the_py_launcher_first(name):
    text = (ROOT / name).read_bytes().decode("ascii")
    assert "\r\n" in text and "\n" not in text.replace("\r\n", "")
    first_py = min(i for i in (text.find("pyw run.py"), text.find("set PY=py\r")) if i >= 0)
    plain = [i for i in (text.find("pythonw run.py"), text.find("python run.py"), text.find("python -")) if i >= 0]
    assert all(first_py < i for i in plain), name


def test_button_is_only_made_inside_an_install(tmp_path):
    main.ensure_launchers(tmp_path)
    assert not any(tmp_path.glob("*.bat"))


def test_new_launchers_are_kept_by_later_updates():
    import updater
    assert set(main.LAUNCHERS) <= set(updater.MANAGED)


def test_unreadable_scan_is_listed_by_its_file_name(home):
    main.fetch()
    names = [r["supplier_name"] for r in view(home)]
    assert "קובץ: 07_scanned.pdf" in names and None not in names
