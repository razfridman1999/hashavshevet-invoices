"""The suppliers' ledger (כרטסת) exported from Hashavshevet: an invoice already in it never
goes in again, and a supplier seen for the first time gets the expense card its invoices
went to before. The layout copies a real export; every number here is invented."""
from __future__ import annotations

import os
import sys
import time
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from invoices import journal, ledger, main  # noqa: E402
from test_workflow import S, fill, home, standard_pass  # noqa: E402,F401  (fixture)

HEAD = [None, None, None, "כותרת", "תנועה", "מנה", 'ס"ת', "ח-ן נגדי", "ת.אסמכ", "ת.ערך", "אסמ'", "אסמ'2",
        "פרטים", "תמחיר", " חובה / זכות (שקל) חובה", "  חובה / זכות (שקל) זכות", "יתרה (שקל)"]


def line(counter, ref1, ref2, amount, credit=True):
    debit, cred = ("", amount) if credit else (amount, "")
    return [None, None, None, "42", "114", "4", "ח8", counter, "25/01/26", "25/01/26", ref1, ref2,
            "01/26 דוגמה", None, debit, cred, amount]


def write_ledger(path: Path, blocks: dict[str, list[list]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["<שם החברה> ח.פ <מספר>"] + [None] * 15 + ["16/9/2026 09:12"])
    ws.append(["כרטסת הנהלת חשבונות"])
    ws.append(["תאריך מ..עד 01/01/2026  >> 31/12/2029"])
    ws.append(HEAD)
    for account, lines in blocks.items():
        ws.append(["ספק לדוגמה", account, "500"])
        ws.append([None, None, "יתרת פתיחה ", None, None, None, None, None, None, None, None, None, None, None,
                   None, None, "0.00"])
        for ln in lines:
            ws.append(ln)
        ws.append(['סה"כ מפתח חשבון', account, "ספק לדוגמה"] + [None] * 12 + ["1,559.96   חובה"])
        ws.append([None] * 15 + ["1,559.96    זכות", "0.00"])
    wb.save(path)
    return path


SAMPLE = {"5916": [line("9302", "0", "50808", "778.80"), line("9302", "0", "50871", "781.16"),
                   line("1100", "17746", None, "1,559.96", credit=False)],
          "6310": [line("9401", "26607409", "26607409", "67,894.34")]}


def test_reads_the_export_layout(tmp_path):
    led = ledger.load(write_ledger(tmp_path / "כרטסת.xlsx", SAMPLE))
    assert led.has("5916", "50808", 778.80)                  # the number sits in אסמ'2, אסמ' is 0
    assert led.has("5916", "3-50871", "781.16")               # her prefix form of the same number
    assert led.has("6310", "26607409", "67,894.34")           # amounts come as text with commas
    assert not led.has("5916", "50808", 778.81)               # same number, other amount
    assert not led.has("6310", "50808", 778.80)               # same number, other supplier
    assert led.expense_accounts("5916") == ["9302"]           # a payment (bank 1100) is never an expense card
    assert led.expense_accounts("9999") == []


def test_most_used_expense_card_first(tmp_path):
    led = ledger.load(write_ledger(tmp_path / "k.xlsx", {"1": [line("A1", "0", "1", "1"), line("B2", "0", "2", "1"),
                                                                 line("B2", "0", "3", "1")]}))
    assert led.expense_accounts("1") == ["B2", "A1"]


def test_a_file_that_is_not_a_ledger_is_refused(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["מפתח חשבון", "שם"])
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(ValueError):
        ledger.load(tmp_path / "x.xlsx")


def test_finds_the_newest_ledger_she_dropped_in_a_day_folder(tmp_path):
    files = tmp_path / "לקוח"
    (files / "10.10.26").mkdir(parents=True)
    (files / "12.10.26").mkdir()
    old = write_ledger(files / "10.10.26" / "כרטסת ספקים.xlsx", SAMPLE)
    new = write_ledger(files / "12.10.26" / "כרטסת ספקים.xlsx", SAMPLE)
    os.utime(old, (time.time() - 3600,) * 2)
    (files / "12.10.26" / "~$כרטסת ספקים.xlsx").write_bytes(b"lock")     # Excel's lock file
    assert ledger.find({"files_folder": str(files)}) == new
    assert ledger.find({"ledger": str(old), "files_folder": str(files)}) == old
    assert ledger.find({"files_folder": ""}) is None


# --- through fetch and export ------------------------------------------------------------
def with_index_and_ledger(home, blocks):
    idx = openpyxl.Workbook()
    idx.active.append(["מפתח חשבון", "שם החשבון", "מס' ע.מורשה"])
    idx.active.append(["40010", 'חשמל השרון בע"מ', S[0]])
    idx.save(home / "index.xlsx")
    cfg = main.load(main.CONFIG, {})
    cfg.update(supplier_index=str(home / "index.xlsx"), ledger=str(write_ledger(home / "כרטסת.xlsx", blocks)))
    main.save(main.CONFIG, cfg)


def electric(home):
    return next(e for e in journal.load(home).values() if e.get("invoice_number") == "20261045"
                and e.get("supplier_name"))


def review_rows(path):
    rows = list(openpyxl.load_workbook(path).active.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def test_new_supplier_gets_its_expense_card_from_the_ledger(home):
    with_index_and_ledger(home, {"40010": [line("81100", "0", "77", "100.00")]})
    rows = review_rows(main.fetch())
    first = next(r for r in rows if r["מס' חשבונית"] == "20261045")
    assert (first["כרטיס ספק"], first["כרטיס הוצאה"]) == ("40010", "81100")
    assert "מהכרטסת" in first["הערות"]


def test_several_expense_cards_means_she_chooses(home):
    with_index_and_ledger(home, {"40010": [line("81100", "0", "77", "1"), line("81200", "0", "78", "1")]})
    first = next(r for r in review_rows(main.fetch()) if r["מס' חשבונית"] == "20261045")
    assert not first["כרטיס הוצאה"]
    assert "ledger_many" in journal.issue_codes({"notes": first["הערות"]})


def test_invoice_already_in_the_ledger_never_goes_in_again(home, capsys):
    with_index_and_ledger(home, {"40010": [line("81100", "0", "20261045", "495.60")]})
    review = main.fetch()
    assert "כבר רשומות בכרטסת" in capsys.readouterr().out
    fill(review, standard_pass)                          # she even marks it "כן"
    out = main.export()
    imported = [r for r in review_rows(out)] if out else []
    assert all(r["אסמכתא 1"] != "20261045" for r in imported)
    assert electric(home)["status"] == journal.REJECTED and electric(home)["todo"] == main.DUPLICATE_DONE


def test_unreadable_ledger_is_said_out_loud_and_work_goes_on(home, capsys):
    wb = openpyxl.Workbook()
    wb.active.append(["לא כרטסת"])
    wb.save(home / "bad.xlsx")
    main.save(main.CONFIG, {**main.load(main.CONFIG, {}), "ledger": str(home / "bad.xlsx")})
    assert main.fetch() is not None
    assert "לא נבדקו מולה כפילויות" in capsys.readouterr().out
