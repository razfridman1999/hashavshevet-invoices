"""An invoice split over several expense cards: one journal line per card, same reference,
the rounding on the largest line. The numbers in the first test are those of a split
invoice that was imported for real, with every number changed: same shape, not hers."""
from __future__ import annotations

import sys
from decimal import Decimal as D
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from invoices import journal, main, scan, sections  # noqa: E402
from invoices.workbook import COL  # noqa: E402
from test_workflow import fill, home, import_rows, standard_pass  # noqa: E402,F401  (fixture)


def amounts(rows):
    return [(r["expense_account"], D(str(r["net"])), D(str(r["vat"])), D(str(r["total"]))) for r in rows]


def test_rounding_goes_on_the_largest_line_and_everything_adds_up():
    rows, problem = sections.parts({"net": 5000.33, "vat": 900.06, "total": 5900.39, "split": "81100=1000; 81200=4000"})
    assert problem == "" and amounts(rows) == [("81100", D("1000.00"), D("180.00"), D("1180.00")),
                                               ("81200", D("4000.33"), D("720.06"), D("4720.39"))]


def test_rounded_vat_shares_never_add_an_agora():
    """Three equal thirds of 0.68 round to 0.23 each (0.69): the largest line absorbs it."""
    rows, _ = sections.parts({"net": 3.75, "vat": 0.68, "total": 4.43, "split": "81=1.25; 82=1.25; 83=1.25"})
    assert sum(D(str(r["vat"])) for r in rows) == D("0.68") and sum(D(str(r["total"])) for r in rows) == D("4.43")


def test_a_credit_note_splits_negative():
    rows, _ = sections.parts({"net": -100.0, "vat": -18.0, "total": -118.0, "split": "81=60; 82=40"})
    assert len(rows) == 2 and amounts(rows) == [("81", D("-60.00"), D("-10.80"), D("-70.80")),
                                                ("82", D("-40.00"), D("-7.20"), D("-47.20"))]


def test_cards_that_do_not_add_up_are_refused():
    row = {"net": 1000.0, "vat": 180.0, "total": 1180.0, "split": "81=600; 82=300"}
    rows, problem = sections.parts(row)
    assert problem and rows == [row]


def test_one_card_or_none_is_the_row_as_it_is():
    row = {"net": 1.0, "vat": 0.18, "total": 1.18, "split": "81100=1"}
    assert sections.parts(row) == ([row], "")
    assert sections.parse("81100 = 2,400.50 ; 81200:10") == [("81100", D("2400.50")), ("81200", D("10.00"))]


def set_split(review, file_start, text):
    wb = openpyxl.load_workbook(review)
    for row in wb.active.iter_rows(min_row=2):
        if Path(row[COL["file"]].value).name.startswith(file_start):
            row[COL["split"]].value = text
    wb.save(review)


def test_export_writes_one_line_per_card_with_the_same_reference(home):
    review = main.fetch()
    fill(review, standard_pass)
    set_split(review, "01", "81100=300; 81200=120")          # the electricity bill: 420.00 before VAT
    lines = [r for r in import_rows(main.export()) if r["אסמכתא 1"] == "20261045"]
    assert [(r["חשבון חובה 1"], r["סכום חובה 1"]) for r in lines] == [("81100", 300.0), ("81200", 120.0)]
    assert round(sum(r["סכום זכות 1"] for r in lines), 2) == 495.60       # the supplier gets the whole invoice
    for r in lines:
        assert round(r["סכום חובה 1"] + r["סכום חובה 2"], 2) == r["סכום זכות 1"]


def test_a_split_that_does_not_add_up_is_held(home, capsys):
    review = main.fetch()
    fill(review, standard_pass)
    set_split(review, "01", "81100=300; 81200=50")
    out = main.export()
    assert all(r["אסמכתא 1"] != "20261045" for r in import_rows(out))
    assert "לא מסתכמים" in capsys.readouterr().out


READ = {"doc_type": "tax_invoice", "supplier_id": "513100065", "invoice_number": "67", "date": "31/07/2026",
        "net": 5000.33, "vat": 900.06, "total": 5900.39, "allocation_number": "300112299", "section": "",
        "sections": [{"card": "81100", "lines": "1", "net": 1000}, {"card": "81200", "lines": "5-15", "net": 4000}],
        "unsure": ["allocation_number"]}


def test_a_stamp_with_two_cards_becomes_a_split_and_unclear_digits_are_named():
    inv = scan.to_invoice(READ, "s.pdf", set())
    assert inv.splits == [("81100", D("1000.00")), ("81200", D("4000.00"))]
    codes = journal.issue_codes({"notes": " | ".join(inv.warnings)})
    assert "scan_unsure" in codes and any("מספר הקצאה" in w for w in inv.warnings)


def test_the_split_reaches_her_sheet(home, monkeypatch):
    monkeypatch.setattr(scan, "claude_path", lambda: "claude.exe")
    monkeypatch.setattr(scan, "read_data", lambda pdf, own, model="sonnet": dict(READ))
    main.save(main.CONFIG, {**main.load(main.CONFIG, {}), "scan_reader": "claude"})
    rows = list(openpyxl.load_workbook(main.fetch()).active.iter_rows(values_only=True))
    row = next(dict(zip(rows[0], r)) for r in rows[1:] if "07_scanned" in str(r[COL["file"]]))
    assert row["פיצול לסעיפים"] == "81100=1000.00; 81200=4000.00" and row["כרטיס הוצאה"] == "81200"
    assert "split" in journal.issue_codes({"notes": row["הערות"]})
