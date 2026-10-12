"""The rules from the bookkeeper's real rounds (Hashavshevet desktop): פרטים "MM/YY ...",
אסמכתא without prefix, ס.ת per company and empty without VAT, מספר עמ / הקצאה,
חשבון עסקה never entered, supplier cards from Hashavshevet's own index."""
from __future__ import annotations

import importlib.util
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "demo"))

import make_demo_invoices as demo  # noqa: E402
from invoices import main, supplier_index  # noqa: E402
from invoices.extract import _extract_text, journal_reference  # noqa: E402
from invoices.workbook import (COL, IMPORT_COLS, REVIEW_COLS, description_of, details_text,  # noqa: E402
                               read_review)

S = demo.S
SUP, OWN = "513100008", "515550002"
BODY = '\nתאריך: 05/10/2026\nסה"כ לפני מע"מ: 100.00\nמע"מ 18%: 18.00\nסה"כ לתשלום: 118.00'
I = {name: i for i, name in enumerate(IMPORT_COLS)}


# --- אסמכתא ----------------------------------------------------------------------------------
@pytest.mark.parametrize("raw, ref", [("3-1901938", "1901938"), ("01/004687", "4687"), ("20261045", "20261045"),
                                      ("INV-0012", "12"), ("1234567890123", "567890123"), ("", "")])
def test_journal_reference(raw, ref):
    assert journal_reference(raw) == ref


@pytest.mark.parametrize("title, number", [("חשבונית מס 3-1901938", "3-1901938"),
                                           ("חשבונית מס 01/004687", "01/004687"),
                                           ("חשבונית מס 05/10/2026", "")])
def test_number_with_series_prefix(title, number):
    inv = _extract_text(f'ספק בע"מ\nח.פ {SUP}\n{title}' + BODY, "x.pdf", {OWN})
    assert inv.invoice_number == number


# --- חשבון עסקה ------------------------------------------------------------------------------
def test_deal_only_document_is_marked():
    assert _extract_text(f'ספק בע"מ\nח.פ {SUP}\nחשבון עסקה 55' + BODY, "x.pdf", {OWN}).deal_only
    assert not _extract_text(f'ספק בע"מ\nח.פ {SUP}\nחשבונית מס 55\nבמקום חשבון עסקה 12' + BODY,
                             "x.pdf", {OWN}).deal_only


# --- פרטים -----------------------------------------------------------------------------------
def test_details_month_first_two_digits():
    assert details_text(date(2026, 9, 3), "ניקיון משרד", "ברק שירותים") == "09/26 ניקיון משרד ברק שירותים"
    assert details_text(date(2026, 9, 3), "", "ברק שירותים") == "09/26 ברק שירותים"
    assert description_of("09/26 ניקיון משרד ברק שירותים", "ברק שירותים") == "ניקיון משרד"


# --- the full cycle with the new columns -------------------------------------------------------
@pytest.fixture
def home(tmp_path, monkeypatch):
    demo.OUT = tmp_path / "mail"
    demo.main()
    h = tmp_path / "home"
    h.mkdir()
    for name, value in (("HOME", h), ("CONFIG", h / "config.json"),
                        ("SUPPLIERS", h / "suppliers.json"), ("STATE", h / "state.json")):
        monkeypatch.setattr(main, name, value)
    main.save(main.CONFIG, {**main.DEFAULT_CONFIG, "source": "folder", "test_folder": str(demo.OUT),
                            "own_company_ids": [demo.OWN_ID], "vat_account": "77100", "ttype": "ח"})
    return h


def review_rows(path):
    return {Path(r["file"]).name: r for r in read_review(path)}


def approve(path, accounts, details=None):
    wb = openpyxl.load_workbook(path)
    for row in wb.active.iter_rows(min_row=2):
        name = Path(row[COL["file"]].value).name
        if name in accounts:
            row[COL["approve"]].value = "כן"
            row[COL["supplier_account"]].value, row[COL["expense_account"]].value = accounts[name]
            if details and name in details:
                row[COL["details"]].value = details[name]
    wb.save(path)


def test_import_columns_ttype_vat_id_allocation(home):
    review = main.fetch()
    approve(review, {"01_electric.pdf": ("40010", "81100"), "03_exempt_dealer.pdf": ("40012", "81300")})
    rows = list(openpyxl.load_workbook(main.export()).active.iter_rows(min_row=2, values_only=True))
    electric = next(r for r in rows if r[I["אסמכתא 1"]] == "20261045")
    exempt = next(r for r in rows if r[I["אסמכתא 1"]] == "112")
    assert electric[I["סוג תנועה"]] == "ח" and electric[I["אסמכתא 2"]] == "20261045"
    assert electric[I["מספר עמ"]] == S[0] and electric[I["מספר הקצאה"]] == "300112233"
    assert electric[I["פרטים"]] == '10/26 חשמל השרון בע"מ'
    assert not exempt[I["סוג תנועה"]]            # no VAT -> ס.ת empty
    assert not exempt[I["חשבון חובה 2"]]


def test_her_details_wording_is_learned_for_next_month(home):
    approve(main.fetch(), {"01_electric.pdf": ("40010", "81100")},
            {"01_electric.pdf": '10/26 חשמל משרד חשמל השרון בע"מ'})
    main.export()
    demo.page(demo.simple('חשמל השרון בע"מ', S[0], "חשבונית מס 20261187", "05/11/2026",
                          "455.00", "81.90", "536.90"), "09_electric_nov.pdf")
    nov = review_rows(main.fetch())["09_electric_nov.pdf"]
    assert nov["details"] == '11/26 חשמל משרד חשמל השרון בע"מ'


def test_credit_note_details_say_zikui(home):
    r = review_rows(main.fetch())["04_credit_note.pdf"]
    assert r["details"].startswith("10/26 זיכוי")


def test_deal_only_is_shown_but_never_exported(home):
    demo.page(f'<h2>שיפוצים בע"מ</h2><p>עוסק מורשה {S[6]}</p><h3>חשבון עסקה 801</h3><p>תאריך: 04/10/2026</p>'
              '<p>סה"כ לפני מע"מ: 100.00</p><p>מע"מ 18%: 18.00</p><p>סה"כ לתשלום: 118.00</p>', "10_deal.pdf")
    review = main.fetch()
    deal = review_rows(review)["10_deal.pdf"]
    assert deal["approve"] == "לא" and "חשבון עסקה" in deal["notes"]
    approve(review, {"01_electric.pdf": ("40010", "81100")})
    rows = list(openpyxl.load_workbook(main.export()).active.iter_rows(min_row=2, values_only=True))
    assert all(r[I["אסמכתא 1"]] != "801" for r in rows)
    assert "10_deal.pdf" not in review_rows(main.fetch() or review)   # rejected for good


# --- supplier index ------------------------------------------------------------------------
def index_rows():
    """Shaped like a real Hashavshevet export: title lines, then the header row."""
    return [[f'חברה לדוגמה בע"מ ח.פ {OWN}'] + [None] * 8, ["אינדקס ספקים"] + [None] * 8,
            [None, None, None, None, "קוד מיון", "מפתח חשבון", "שם החשבון", "חשבון ראשי", "מס' ע.מורשה"],
            ["קוד מיון", 500, "ספקים", 500, None, None, None, None, None],
            [None, None, None, None, 500, 6310, "חשמל לדוגמה", 500, S[0]],
            [None, None, None, None, 500, "5575", "משרדית לדוגמה", 500, int(S[1])],
            [None, None, None, None, 500, "6400", "כפול א", 500, S[2]],
            [None, None, None, None, 500, "6401", "כפול ב", 500, S[2]]]


def test_index_parse_matches_by_vat_id():
    idx = supplier_index.parse(index_rows())
    assert idx[S[0]] == [{"supplier_account": "6310", "name": "חשמל לדוגמה"}]
    assert idx[S[1]][0]["supplier_account"] == "5575"
    assert len(idx[S[2]]) == 2


def test_new_supplier_gets_card_from_index(home, tmp_path):
    wb = openpyxl.Workbook()
    for r in index_rows():
        wb.active.append(r)
    wb.save(tmp_path / "index.xlsx")
    cfg = main.load(main.CONFIG, {})
    main.save(main.CONFIG, {**cfg, "supplier_index": str(tmp_path / "index.xlsx")})
    rows = review_rows(main.fetch())
    electric, office, exempt = rows["01_electric.pdf"], rows["02_office_table.pdf"], rows["03_exempt_dealer.pdf"]
    assert electric["supplier_account"] == "6310" and electric["supplier_name"] == "חשמל לדוגמה"
    assert not electric["expense_account"] and "מהאינדקס" in electric["notes"]
    assert office["supplier_account"] == "5575"
    assert not exempt["supplier_account"] and "כמה כרטיסי ספק" in exempt["notes"]


# --- an unfinished sheet from the previous version still reads correctly ---------------------
def test_old_version_sheet_is_read_by_header(tmp_path):
    old_cols = [c for c in REVIEW_COLS if c[0] != "details"]
    wb = openpyxl.Workbook()
    wb.active.append([t for _, t in old_cols])
    wb.active.append(["כן", "ספק", SUP, "40010", "81100", "77", None, 100, 18, 118, "", "", "x.pdf", "k1", ""])
    wb.save(tmp_path / "old.xlsx")
    r = read_review(tmp_path / "old.xlsx")[0]
    assert (r["file"], r["key"], r["details"], r["expense_account"]) == ("x.pdf", "k1", "", "81100")


# --- updater: a local copy ahead of the release is never rolled back ----------------------------
def test_updater_never_downgrades():
    spec = importlib.util.spec_from_file_location("updater_check", ROOT / "updater.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    assert mod.is_newer("2026.10.12.3", "2026.10.12.2")
    assert mod.is_newer("2026.10.13.1", "2026.10.12.9")
    assert not mod.is_newer("2026.10.12.2", "2026.10.12.3")
    assert not mod.is_newer("2026.10.12.3", "2026.10.12.3")
    assert mod.is_newer("1.0", "")
