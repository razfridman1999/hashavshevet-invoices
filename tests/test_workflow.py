"""The full fetch -> review -> export cycle, as the bookkeeper runs it."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "demo"))

import make_demo_invoices as demo  # noqa: E402
from invoices import main  # noqa: E402

S = demo.S
ACCOUNTS = {S[0]: (40010, 81100), S[1]: (40011, 81200), S[2]: (40012, 81300),
            S[3]: (40013, 81400), S[4]: (40014, 81500)}


@pytest.fixture
def home(tmp_path, monkeypatch):
    demo.OUT = tmp_path / "mail"
    demo.main()
    h = tmp_path / "home"
    h.mkdir()
    for name, value in (("HOME", h), ("CONFIG", h / "config.json"),
                        ("SUPPLIERS", h / "suppliers.json"), ("STATE", h / "state.json")):
        monkeypatch.setattr(main, name, value)
    main.save(main.CONFIG, {**main.DEFAULT_CONFIG, "source": "folder", "test_folder": str(demo.OUT),
                            "own_company_ids": [demo.OWN_ID], "vat_account": "77100", "ttype": "ח"})
    return h


def fill(review: Path, decide) -> None:
    """Play the bookkeeper: decide(row_cells, file_name) edits a row in place."""
    wb = openpyxl.load_workbook(review)
    for row in wb.active.iter_rows(min_row=2):
        decide(row, Path(row[13].value).name)
    wb.save(review)


def standard_pass(row, name):
    if name.startswith(("06", "08")):
        row[0].value = "לא"
    elif not name.startswith("07"):
        row[0].value = "כן"
        row[3].value, row[4].value = ACCOUNTS[row[2].value]


def import_rows(path: Path) -> list[tuple]:
    return list(openpyxl.load_workbook(path).active.iter_rows(min_row=2, values_only=True))


def test_full_cycle_balances_and_learns(home):
    review = main.fetch()
    fill(review, standard_pass)
    out = main.export()
    rows = import_rows(out)
    assert len(rows) == 5
    for r in rows:  # every journal line balances: debits == credits
        assert round((r[7] or 0) + (r[9] or 0), 2) == round((r[11] or 0) + (r[13] or 0), 2)
    credit = next(r for r in rows if "זיכוי" in r[5])
    assert (credit[6], credit[10], credit[12]) == ("40013", "81400", "77100")  # supplier debited
    assert json.loads((home / "suppliers.json").read_text(encoding="utf-8"))[S[0]]["supplier_account"] == "40010"
    # second run: only the scan, which she has not handled yet, comes back
    review2 = main.fetch()
    names = [Path(r[13]).name for r in openpyxl.load_workbook(review2).active.iter_rows(min_row=2, values_only=True)]
    assert names == ["07_scanned.pdf"]


def test_known_supplier_comes_back_ready(home):
    fill(main.fetch(), standard_pass)
    main.export()
    demo.page(demo.simple('חשמל השרון בע"מ', S[0], "חשבונית מס 20261187", "05/11/2026",
                          "455.00", "81.90", "536.90"), "09_electric_nov.pdf")
    rows = {Path(r[13]).name: r for r in
            openpyxl.load_workbook(main.fetch()).active.iter_rows(min_row=2, values_only=True)}
    nov = rows["09_electric_nov.pdf"]
    assert (nov[0], nov[3], nov[4], nov[12]) == ("כן", "40010", "81100", None)


def test_duplicate_approved_twice_enters_once(home):
    def approve_everything(row, name):
        standard_pass(row, name)
        if name.startswith("08"):          # she approves the forwarded copy too
            row[0].value = "כן"
            row[3].value, row[4].value = ACCOUNTS[S[0]]
    fill(main.fetch(), approve_everything)
    rows = import_rows(main.export())
    assert sum(1 for r in rows if r[1] == "20261045") == 1


def test_hand_edited_amount_that_does_not_balance_is_held_back(home):
    def bad_edit(row, name):
        standard_pass(row, name)
        if name.startswith("01"):
            row[9].value = 999.99          # total no longer equals net + VAT
    fill(main.fetch(), bad_edit)
    rows = import_rows(main.export())
    assert all(r[1] != "20261045" for r in rows)
    assert len(rows) == 4


def test_unfinished_sheet_survives_a_second_fetch(home):
    first = main.fetch()
    def half_done(row, name):
        if name.startswith("01"):
            row[3].value, row[4].value = ACCOUNTS[S[0]]
    fill(first, half_done)
    first_mtime = first.stat().st_mtime
    import os, time
    os.utime(first, (first_mtime, first_mtime - 120))   # older file, as if made earlier today
    time.sleep(0.01)
    second = main.fetch()
    rows = {Path(r[13]).name: r for r in openpyxl.load_workbook(second).active.iter_rows(min_row=2, values_only=True)}
    assert (rows["01_electric.pdf"][3], rows["01_electric.pdf"][4]) == ("40010", "81100")


# --- review findings 3, 7, 8, 9: state across runs ------------------------------------------
def test_exporting_the_same_sheet_twice_enters_nothing_twice(home):
    review = main.fetch()
    fill(review, standard_pass)
    first = import_rows(main.export(review))
    second = main.export(review)
    assert len(first) == 5
    assert second is None or import_rows(second) == []


def test_blank_invoice_number_is_held_back(home):
    def blank(row, name):
        standard_pass(row, name)
        if name.startswith("01"):
            row[5].value = None
    fill(main.fetch(), blank)
    assert all("חשמל" not in (r[5] or "") for r in import_rows(main.export()))


def test_corrupt_pdf_does_not_stop_the_others(home):
    (demo.OUT / "00_corrupt.pdf").write_bytes(b"%PDF-1.4 broken")
    rows = list(openpyxl.load_workbook(main.fetch()).active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 9
    assert any(Path(r[13]).name == "00_corrupt.pdf" and r[12] for r in rows)


def test_pending_row_survives_when_its_mail_is_no_longer_fetched(home):
    first = main.fetch()
    def half(row, name):
        if name.startswith("02"):
            row[3].value, row[4].value = ACCOUNTS[S[1]]
    fill(first, half)
    (demo.OUT / "02_office_table.pdf").unlink()        # aged out of days_back / moved
    import os
    os.utime(first, (first.stat().st_mtime - 120,) * 2)
    rows = {Path(r[13]).name: r for r in
            openpyxl.load_workbook(main.fetch()).active.iter_rows(min_row=2, values_only=True)}
    assert rows["02_office_table.pdf"][3] == "40011"


def test_hand_added_row_exports_once_and_does_not_crash(home):
    review = main.fetch()
    wb = openpyxl.load_workbook(review)
    ws = wb.active
    ws.append(["כן", "מוסך הגליל", S[6], "40020", "81600", "4410", None, 600, 108, 708, None, "10/26 תיקון רכב מוסך הגליל", None, None, None, None])
    ws.cell(ws.max_row, 7).value = __import__("datetime").datetime(2026, 10, 2)
    wb.save(review)
    first = import_rows(main.export(review))
    assert sum(1 for r in first if r[1] == "4410") == 1
    second = main.export(review)
    assert second is None or all(r[1] != "4410" for r in import_rows(second))


@pytest.mark.parametrize("suffix", [".xlsx", ".xls"])
def test_hashavshevet_can_read_it_over_odbc(home, suffix):
    """Read the file the way Hashavshevet does: ODBC, table = the named range."""
    pyodbc = pytest.importorskip("pyodbc")
    if not any("Excel" in d for d in pyodbc.drivers()):
        pytest.skip("no Excel ODBC driver")
    fill(main.fetch(), standard_pass)
    out = main.export().with_suffix(suffix)
    if not out.exists():
        pytest.skip("Excel not installed, no .xls copy")
    cn = pyodbc.connect("DRIVER={Microsoft Excel Driver (*.xls, *.xlsx, *.xlsm, *.xlsb)};"
                        f"DBQ={out};ReadOnly=1", autocommit=True)
    cur = cn.cursor()
    cur.execute("SELECT * FROM invoices")
    cols = [d[0] for d in cur.description]
    rows = cur.fetchall()
    cn.close()
    assert len(rows) == 5
    assert isinstance(rows[0][cols.index("חשבון זכות 1")], str)   # account keys must stay text
    assert isinstance(rows[0][cols.index("סכום זכות 1")], float)


@pytest.mark.parametrize("suffix", [".xlsx", ".xls"])
def test_odbc_credit_note_after_many_plain_rows(home, suffix):
    """Review finding 11: ODBC guesses column types from the first rows. A credit note far
    down the sheet must still have its second credit (the VAT) read, not NULL."""
    pyodbc = pytest.importorskip("pyodbc")
    for i in range(10):
        demo.page(demo.simple('חשמל השרון בע"מ', S[0], f"חשבונית מס 3000{i}", "05/10/2026",
                              "100.00", "18.00", "118.00"), f"00_{i:02}_electric.pdf")
    fill(main.fetch(), standard_pass)
    out = main.export().with_suffix(suffix)
    if not out.exists():
        pytest.skip("no .xls copy")
    cn = pyodbc.connect("DRIVER={Microsoft Excel Driver (*.xls, *.xlsx, *.xlsm, *.xlsb)};"
                        f"DBQ={out};ReadOnly=1", autocommit=True)
    cur = cn.cursor()
    cur.execute("SELECT * FROM invoices")
    cols = [d[0] for d in cur.description]
    rows = cur.fetchall()
    cn.close()
    credit = [r for r in rows if "זיכוי" in r[cols.index("פרטים")]]
    assert len(credit) == 1
    assert credit[0][cols.index("חשבון זכות 2")] == "77100"
    assert credit[0][cols.index("סכום זכות 2")] == 36.0
    for r in rows:  # every row balances as ODBC sees it
        d = (r[cols.index("סכום חובה 1")] or 0) + (r[cols.index("סכום חובה 2")] or 0)
        c = (r[cols.index("סכום זכות 1")] or 0) + (r[cols.index("סכום זכות 2")] or 0)
        assert round(d, 2) == round(c, 2)
